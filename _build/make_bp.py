# -*- coding: utf-8 -*-
"""Build the two business-plan documents: 设计纲要 + 简略样板.

Two-pass build (page numbers first, then a populated TOC), 宋体 throughout.
"""

import json
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION

import docgen as G
from fontfix import fix_fonts_in_docx
from bp_outline import OPS_OUTLINE
from bp_sample import OPS_SAMPLE

OUT_DIR = os.path.dirname(HERE)
PDF_DIR = os.path.join(HERE, 'pdf')
PROJECT = '储粮安全物联网监控平台'


def _find_soffice():
    for c in (r'C:\Program Files\LibreOffice\program\soffice.exe',
              r'C:\Program Files (x86)\LibreOffice\program\soffice.exe'):
        if os.path.exists(c):
            return c
    return 'soffice'


SOFFICE = _find_soffice()

DOCS = [
    {
        'name': '%s-商业计划书设计纲要' % PROJECT,
        'title': '商业计划书设计纲要',
        'banner': '参 赛 材 料 编 制 系 列',
        'cover_titles': ['商业计划书', '设 计 纲 要'],
        'cover_sub': '—— 逐章写法 · 篇幅预算 · 图表清单 · 自检清单 ——',
        'header': '商业计划书设计纲要',
        'meta': [
            ('适用项目', PROJECT),
            ('文档用途', '商业计划书写作的施工图'),
            ('配套文档', '《项目简介与BP-PPT编制指南》《商业计划书简略样板》'),
            ('团队规模', '7 人'),
        ],
        'ops': OPS_OUTLINE,
    },
    {
        'name': '%s-商业计划书简略样板' % PROJECT,
        'title': '商业计划书（简略样板）',
        'banner': '参 赛 材 料 编 制 系 列',
        'cover_titles': ['商业计划书', '（简 略 样 板）'],
        'cover_sub': '—— %s ——' % PROJECT,
        'header': '商业计划书（简略样板）',
        'meta': [
            ('项目名称', PROJECT),
            ('文档性质', '压缩样板，用于演示写法，不可直接提交'),
            ('团队名称', '【待填】'),
            ('所在学校', '【待填】'),
        ],
        'ops': OPS_SAMPLE,
    },
]


def build(spec, entries, out_name):
    doc = G.new_document()

    G.cover(doc,
            titles=spec['cover_titles'],
            meta=spec['meta'],
            banner_line=spec['banner'],
            subtitle=spec['cover_sub'])
    cover_sec = doc.sections[0]
    G.clear_header(cover_sec)
    G.clear_footer(cover_sec)

    doc.add_section(WD_SECTION.NEW_PAGE)
    toc_sec = doc.sections[1]
    G.clear_header(toc_sec)
    G.clear_footer(toc_sec)
    G.toc(doc, entries, line_exact=14)

    doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec = doc.sections[2]
    G.header_text(body_sec, spec['header'])
    G.footer_page_number(body_sec)
    G.set_page_numbering(body_sec, start=1, fmt='decimal')

    G.render(doc, spec['ops'])
    G.enable_update_fields(doc)

    props = doc.core_properties
    props.title = spec['title'] + '（%s）' % PROJECT
    props.subject = '挑战杯中国大学生创业计划竞赛 参赛材料'
    props.author = '【待填：团队名称】'

    path = os.path.join(OUT_DIR, out_name + '.docx')
    try:
        doc.save(path)
    except PermissionError:
        out_name = out_name + '(新版本)'
        path = os.path.join(OUT_DIR, out_name + '.docx')
        doc.save(path)
        print('!! 原文档被占用（可能已在 WPS / Word 中打开），改存为 %s.docx' % out_name)
    fix_fonts_in_docx(path)
    return out_name


def to_pdf(name):
    os.makedirs(PDF_DIR, exist_ok=True)
    subprocess.run([SOFFICE, '--headless', '--convert-to', 'pdf',
                    '--outdir', PDF_DIR, os.path.join(OUT_DIR, name + '.docx')],
                   check=True, timeout=900)
    return os.path.join(PDF_DIR, name + '.pdf')


def run(spec):
    expected = [(1 if k == 'h1' else 2, v)
                for k, v in spec['ops'] if k in ('h1', 'h2')]

    name = build(spec, None, spec['name'])
    pdf1 = to_pdf(name)

    json_path = os.path.join(HERE, '_toc_%s.json' % spec['name'])
    subprocess.run([sys.executable, os.path.join(HERE, 'extract_toc.py'),
                    pdf1, json_path], check=True, cwd=HERE)
    rows = json.load(open(json_path, encoding='utf-8'))

    if len(rows) != len(expected):
        print('!! %s: heading count mismatch extracted=%d expected=%d'
              % (spec['name'], len(rows), len(expected)))
        print('   extracted:', [r['text'] for r in rows])

    entries = []
    for i, (level, text) in enumerate(expected):
        page = rows[i]['page'] if i < len(rows) else 0
        entries.append((level, text.replace('\u3000', ' ').strip(), page))

    name = build(spec, entries, spec['name'])
    pdf2 = to_pdf(name)

    try:
        shutil.copyfile(pdf2, os.path.join(OUT_DIR, name + '(预览版).pdf'))
    except PermissionError:
        print('!! 预览 PDF 被占用，未更新')

    d = pymupdf.open(pdf2)
    docx = Document(os.path.join(OUT_DIR, name + '.docx'))
    print('== %s' % name)
    print('   pages=%d  paragraphs=%d  tables=%d  toc_entries=%d'
          % (d.page_count, len(docx.paragraphs), len(docx.tables), len(entries)))
    return name


def main():
    for spec in DOCS:
        run(spec)


if __name__ == '__main__':
    main()
